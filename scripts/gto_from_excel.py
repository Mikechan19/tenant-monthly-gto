#!/usr/bin/env python3
"""Generate monthly GTO actuals from the leasing team's Excel workings — no hand-typed figures.

Everything that identifies a tenant (workbook and sheet names, tenant codes, properties) lives in a
local config file, private/gto_sources.json, which is gitignored like every output of this script.

Usage:
  python3 scripts/gto_from_excel.py [EXCEL_DIR]        (or env GTO_EXCEL_DIR)
  config: private/gto_sources.json                      (or env GTO_SOURCES)

Two workbook layouts are understood:

1. "monthly_calc" — one sheet per tenant per month, named <MON>(<CODE>) e.g. "JAN(AB)":
     C1 tenant name · D3 gross sales
     from row 6: B band no · C band label ("0 - 400,000", "400,001 - Above") · D rate · E sales · F GTO
     next row with only E/F filled: totals · a row whose D says "AFTER SST (n%)": F = GTO incl. SST
   plus an optional summary sheet (month in column A from row 5, a Total row) whose columns are
   given per tenant code in the config — read only to document its figures.
2. "outlet_summary" — one sheet for one tenant with several outlets:
     A2 tenant name · row 5 headers: Month | <outlet> (RM) | … | Total (RM) | GTO (RM)
     rows 6–17 months (blank outlet cells = no sales yet)
     a row with "Tenancy Term" header: next row B "dd/mm/yyyy – dd/mm/yyyy", C rent psf, D monthly base rent
     a row with "Percentage of Monthly Gross Sales" header: next rows B band label, C rate

Config (private/gto_sources.json):
  {"year": 2026, "team_id": "<uuid>",
   "sources": [
     {"layout": "monthly_calc", "file": "<workbook>.xlsx",
      "tenants": {"<CODE>": {"property": "<name or null>", "summary_cols": {"gross": 2, "gto": 5}}},
      "summary_sheet": "summary", "summary_total_row": 17},
     {"layout": "outlet_summary", "file": "<workbook>.xlsx", "sheet": "<sheet>", "property": null}]}

Writes (all gitignored):
  private/<year>_gto_actuals.sql   idempotent seed: tenants + terms upserted, monthly sales inserted once
  private/excel_gto_<year>.json    every Excel figure, for __tests__/gto.test.ts and e2e/gto-flow.mjs
  <EXCEL_DIR>/gto_actuals_<year>.csv the same sales rows in the app's CSV-import format
"""
import csv
import json
import os
import re
import sys
from datetime import datetime
from decimal import ROUND_HALF_UP, Decimal
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parent.parent
PRIVATE = ROOT / "private"
DEMO_TEAM = "00000000-0000-4000-8000-00000000d0e0"
MONTHS = ["JAN", "FEB", "MAR", "APR", "MAY", "JUN", "JUL", "AUG", "SEP", "OCT", "NOV", "DEC"]


def clean_name(s):
    return re.sub(r"\s+", " ", str(s)).strip()


def dec(v, places=10):
    """Excel stores IEEE doubles; round to 10 dp to drop binary noise (x.074999999997 → x.075)."""
    return Decimal(repr(float(v))).quantize(Decimal(1).scaleb(-places), rounding=ROUND_HALF_UP).normalize()


def dstr(v):
    """Exact decimal in plain notation, e.g. '3000', '18750.045'."""
    return format(dec(v), "f")


def cents(v):
    return str(dec(v).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP))


def num(v):
    """Plain JSON/SQL number from an Excel value (no float noise)."""
    d = dec(v, 2)
    return int(d) if d == d.to_integral() else float(d)


def parse_band_bound(label):
    """'0 - 400,000' → 400000 ; '400,001 - Above' / '… – above' → None."""
    hi = re.split(r"\s*[-–]\s*", str(label).strip())[-1]
    if hi.lower().startswith("above"):
        return None
    return num(hi.replace(",", ""))


def read_monthly_calc(src, excel_dir):
    wb = openpyxl.load_workbook(excel_dir / src["file"], data_only=True)
    codes = src["tenants"]
    pat = re.compile(r"([A-Z]{3})\((" + "|".join(map(re.escape, codes)) + r")\)")
    tenants = {}
    for ws in wb.worksheets:
        m = pat.fullmatch(ws.title)
        if not m:
            continue
        mon, code = m.groups()
        month = MONTHS.index(mon) + 1
        bands, total, after_sst, sst_pct = [], None, None, 0.0
        for row in ws.iter_rows(min_row=6):
            b, c, d, e, f = (row[i].value for i in range(1, 6))
            if isinstance(b, (int, float)) and c is not None:
                bands.append({"label": str(c).strip(), "rate": float(d), "sales": num(e),
                              "excel_exact": dstr(f), "excel_cents": cents(f)})
            elif b is None and c is None and e is not None and f is not None and total is None:
                total = f
            if isinstance(d, str) and "AFTER SST" in d.upper():
                after_sst = f
                sst_pct = float(re.search(r"([\d.]+)\s*%", d).group(1))
        if total is None:
            raise SystemExit(f"{src['file']} / {ws.title}: no totals row found")
        if after_sst is None:
            after_sst = total
        t = tenants.setdefault(code, {"code": code, "name": clean_name(ws["C1"].value),
                                      "property": codes[code].get("property"), "months": {}})
        tiers = [{"up_to": parse_band_bound(bd["label"]), "rate": bd["rate"]} for bd in bands]
        if "tiers" in t and t["tiers"] != tiers:
            raise SystemExit(f"{ws.title}: bands differ from earlier months: {tiers} vs {t['tiers']}")
        t["tiers"], t["sst_pct"] = tiers, sst_pct
        t["months"][month] = {
            "gross": num(ws["D3"].value), "bands": bands,
            "gto_exact": dstr(total), "gto_cents": cents(total),
            "after_sst_exact": dstr(after_sst), "after_sst_cents": cents(after_sst),
            "sheet": ws.title,
        }
    summary = {}
    if src.get("summary_sheet") in wb.sheetnames:
        s = wb[src["summary_sheet"]]
        for code, cfg in codes.items():
            cols = cfg.get("summary_cols")
            if not cols:
                continue
            out = summary.setdefault(code, {})
            for r in range(5, s.max_row + 1):
                mdate = s.cell(r, 1).value
                if isinstance(mdate, datetime):
                    out[mdate.month] = {"gross": s.cell(r, cols["gross"]).value, "gto": s.cell(r, cols["gto"]).value}
            tr = src.get("summary_total_row")
            if tr:
                out["total"] = {"gross": s.cell(tr, cols["gross"]).value, "gto": s.cell(tr, cols["gto"]).value}
    missing = set(codes) - set(tenants)
    if missing:
        raise SystemExit(f"{src['file']}: no monthly sheets for {sorted(missing)}")
    return [tenants[c] for c in codes], summary


def read_outlet_summary(src, excel_dir):
    ws = openpyxl.load_workbook(excel_dir / src["file"], data_only=True)[src["sheet"]]
    name = clean_name(ws["A2"].value)
    headers = [ws.cell(5, c).value for c in range(1, ws.max_column + 1)]
    outlet_cols = [(i + 1, re.sub(r"\s*\(RM\)\s*$", "", h).strip()) for i, h in enumerate(headers)
                   if i > 0 and isinstance(h, str) and not re.match(r"\s*(total|gto)\b", h, re.I)]
    total_col = next(i + 1 for i, h in enumerate(headers) if isinstance(h, str) and re.match(r"\s*total", h, re.I))
    gto_col = next(i + 1 for i, h in enumerate(headers) if isinstance(h, str) and re.match(r"\s*gto", h, re.I))
    months = {}
    for r in range(6, 18):
        mdate = ws.cell(r, 1).value
        if not isinstance(mdate, datetime):
            continue
        vals = {o: ws.cell(r, c).value for c, o in outlet_cols}
        if all(v is None for v in vals.values()):
            continue  # no sales reported yet (the Total/GTO cells are formulas over blanks)
        g = ws.cell(r, gto_col).value or 0
        months[mdate.month] = {"outlets": {o: num(v or 0) for o, v in vals.items()},
                               "gross": num(ws.cell(r, total_col).value),
                               "gto_exact": dstr(g), "gto_cents": cents(g)}
    lease_start = lease_end = base = psf = None
    tiers = []
    for r in range(1, ws.max_row + 1):
        hdr = " ".join(str(ws.cell(r, c).value or "") for c in range(1, 5))
        if re.search(r"tenancy term", hdr, re.I):
            term = str(ws.cell(r + 1, 2).value)
            lease_start, lease_end = [datetime.strptime(x.strip(), "%d/%m/%Y").date().isoformat() for x in re.split(r"[–-]", term)]
            psf, base = num(ws.cell(r + 1, 3).value), num(ws.cell(r + 1, 4).value)
        if re.search(r"percentage of monthly gross sales", hdr, re.I):
            rr = r + 1
            while ws.cell(rr, 2).value is not None:
                tiers.append({"up_to": parse_band_bound(ws.cell(rr, 2).value), "rate": float(ws.cell(rr, 3).value)})
                rr += 1
    if not tiers:
        raise SystemExit(f"{src['file']}: no turnover-rent bands found")
    return {"code": src["sheet"], "name": name, "property": src.get("property"), "outlets": [o for _, o in outlet_cols],
            "tiers": tiers, "sst_pct": 0, "base_rent_monthly": base, "base_rent_psf": psf,
            "lease_start": lease_start, "lease_end": lease_end, "months": months}


# ── outputs ──────────────────────────────────────────────────────────────────
def sql_str(v):
    return "null" if v is None else "'" + str(v).replace("'", "''") + "'"


def sql_num(v):
    return "null" if v is None else repr(v)


def seed_sql(tenants, year, team):
    out = [
        f"-- {year} monthly GTO actuals — GENERATED locally by scripts/gto_from_excel.py from the leasing",
        "--   team's Excel workings. Contains real tenant data: keep it in private/ (gitignored), never commit.",
        "--",
        "-- Run in the Supabase SQL editor AFTER 0003_gto_terms.sql. Idempotent:",
        "--   * tenants are matched on (team, name); new ones are created, existing ones get their GTO terms",
        "--     updated (fields the Excel doesn't give are left as they are);",
        "--   * a monthly sales row is inserted only if that tenant/month/outlet row doesn't exist yet.",
        "-- Target team: set v_team or v_team_name below.",
        "do $seed$",
        "declare",
        f"  v_team uuid := '{team}';",
        "  v_team_name text := null;  -- e.g. your team's name — if set, overrides v_team",
        "  v_t uuid;",
        "  v_n int;",
        "  v_total int := 0;",
        "begin",
        "  if v_team_name is not null then",
        "    select id into v_team from public.teams where name = v_team_name;",
        "    if v_team is null then raise exception 'No team named %', v_team_name; end if;",
        "  end if;",
        "  if not exists (select 1 from public.teams where id = v_team) then",
        "    raise exception 'Team % not found', v_team;",
        "  end if;",
    ]
    for t in tenants:
        tiers = json.dumps(t["tiers"], separators=(",", ":"))
        base = t.get("base_rent_monthly")
        out += [
            "",
            f"  -- {t['name']}",
            "  v_t := null;",
            f"  select id into v_t from public.tenants where team_id = v_team and name = {sql_str(t['name'])} order by created_at limit 1;",
            "  if v_t is null then",
            "    insert into public.tenants (team_id, name, property, gto_tiers, gto_sst_pct, base_rent_monthly, lease_start, lease_end)",
            f"    values (v_team, {sql_str(t['name'])}, {sql_str(t['property'])}, '{tiers}'::jsonb, {sql_num(t['sst_pct'])},",
            f"            {sql_num(base)}, {sql_str(t.get('lease_start'))}, {sql_str(t.get('lease_end'))})",
            "    returning id into v_t;",
            "  else",
            "    update public.tenants set",
            f"      gto_tiers = '{tiers}'::jsonb, gto_sst_pct = {sql_num(t['sst_pct'])},",
            f"      property = coalesce({sql_str(t['property'])}, property),",
            f"      base_rent_monthly = coalesce({sql_num(base)}, base_rent_monthly),",
            f"      lease_start = coalesce({sql_str(t.get('lease_start'))}::date, lease_start),",
            f"      lease_end = coalesce({sql_str(t.get('lease_end'))}::date, lease_end)",
            "    where id = v_t;",
            "  end if;",
        ]
        rows = []
        for mon, m in sorted(t["months"].items()):
            at = f"{year}-{mon:02d}-01 12:00:00+08"
            if "outlets" in m:
                rows += [(amt, at, o) for o, amt in m["outlets"].items()]
            else:
                rows.append((m["gross"], at, None))
        values = ",\n      ".join(f"({sql_num(a)}::numeric, '{at}'::timestamptz, {sql_str(o)}::text)" for a, at, o in rows)
        out += [
            "  insert into public.sales_records (team_id, tenant_id, period_type, amount, recorded_at, source, outlet)",
            "  select v_team, v_t, 'monthly', x.amount, x.at, 'import', x.outlet",
            f"  from (values\n      {values}\n  ) as x(amount, at, outlet)",
            "  where not exists (",
            "    select 1 from public.sales_records s",
            "    where s.tenant_id = v_t and s.period_type = 'monthly' and s.recorded_at = x.at and s.outlet is not distinct from x.outlet",
            "  );",
            "  get diagnostics v_n = row_count;",
            "  v_total := v_total + v_n;",
            f"  raise notice '{t['name'].replace(chr(39), chr(39) * 2)}: % new monthly sales rows (of {len(rows)})', v_n;",
        ]
    out += ["", "  raise notice 'Done — % sales rows inserted into team %', v_total, v_team;", "end", "$seed$;", ""]
    return "\n".join(out)


def main():
    excel_dir = Path(sys.argv[1] if len(sys.argv) > 1 else os.environ.get("GTO_EXCEL_DIR", "."))
    cfg_path = Path(os.environ.get("GTO_SOURCES", PRIVATE / "gto_sources.json"))
    if not cfg_path.exists():
        raise SystemExit(f"config not found: {cfg_path} (see the docstring for its format)")
    cfg = json.loads(cfg_path.read_text())
    year, team = int(cfg.get("year", 2026)), cfg.get("team_id", DEMO_TEAM)

    tenants, summary = [], {}
    for src in cfg["sources"]:
        if src["layout"] == "monthly_calc":
            ts, s = read_monthly_calc(src, excel_dir)
            tenants += ts
            summary.update(s)
        elif src["layout"] == "outlet_summary":
            tenants.append(read_outlet_summary(src, excel_dir))
        else:
            raise SystemExit(f"unknown layout {src['layout']!r}")

    PRIVATE.mkdir(exist_ok=True)
    (PRIVATE / f"{year}_gto_actuals.sql").write_text(seed_sql(tenants, year, team))
    fixture = {"year": year, "tenants": tenants, "excel_summary_sheet": summary}
    (PRIVATE / f"excel_gto_{year}.json").write_text(json.dumps(fixture, indent=1, default=str) + "\n")

    n = 0
    with open(excel_dir / f"gto_actuals_{year}.csv", "w", newline="") as f:
        w = csv.writer(f)
        w.writerow(["tenant", "date", "amount", "period_type", "outlet"])
        for t in tenants:
            for mon, m in sorted(t["months"].items()):
                for o, amt in (m["outlets"].items() if "outlets" in m else [(None, m["gross"])]):
                    w.writerow([t["name"], f"{year}-{mon:02d}-01", amt, "monthly", o or ""])
                    n += 1
    print(f"{len(tenants)} tenants, {n} monthly sales rows")
    print(f"wrote {PRIVATE / f'{year}_gto_actuals.sql'}, {PRIVATE / f'excel_gto_{year}.json'}, {excel_dir / f'gto_actuals_{year}.csv'}")


if __name__ == "__main__":
    main()
